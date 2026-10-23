#!/usr/bin/env python3
"""Täglicher Einkauf: Bedarfsliste gegen die ITscope Produkt-API prüfen und
pro Artikel den günstigsten Lieferanten mit ausreichendem Lagerbestand wählen.

Aufruf:
    python einkauf.py                      # alle Listen im Eingangsordner verarbeiten
    python einkauf.py pfad\\zur\\Liste.xls  # genau diese Datei
    python einkauf.py Liste.xls --mock beispiel\\api_antworten   # Test ohne API
"""
from __future__ import annotations

import argparse
import contextlib
import io
import json
import logging
import os
import re
import shutil
import sys
import time
import tomllib
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from urllib.parse import quote

import pandas as pd
import requests
from dotenv import load_dotenv
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BASIS = Path(__file__).resolve().parent
log = logging.getLogger("einkauf")

# Status je Zeile
OK = "OK"
TEILBESTAND = "TEILBESTAND"
KEIN_BESTAND = "KEIN BESTAND"
KEIN_ANGEBOT = "KEIN ANGEBOT (meine Lieferanten)"
NICHT_GEFUNDEN = "NICHT GEFUNDEN"
FEHLER = "FEHLER"

ROT = PatternFill("solid", fgColor="FFC7CE")
ORANGE = PatternFill("solid", fgColor="FFE0B2")
GELB = PatternFill("solid", fgColor="FFF2CC")
KOPF = PatternFill("solid", fgColor="D9E1F2")
FARBE = {NICHT_GEFUNDEN: ROT, FEHLER: ROT, KEIN_ANGEBOT: ORANGE,
         TEILBESTAND: GELB, KEIN_BESTAND: GELB}


# --------------------------------------------------------------------------- #
# Hilfsfunktionen
# --------------------------------------------------------------------------- #
def hole(obj, pfade):
    """Ersten vorhandenen Wert aus einer Liste Punkt-getrennter Pfade holen."""
    for pfad in pfade:
        wert = obj
        for teil in pfad.split("."):
            if isinstance(wert, dict) and teil in wert:
                wert = wert[teil]
            else:
                wert = None
                break
        if wert is not None:
            return wert
    return None


def als_liste(wert):
    if wert is None:
        return []
    return wert if isinstance(wert, list) else [wert]


def zahl(wert):
    """'1.234,56', '12.5', 12, {'value': 3} -> float; sonst None."""
    if isinstance(wert, dict):
        wert = hole(wert, ["value", "amount", "net"])
    if wert is None or isinstance(wert, bool):
        return None
    if isinstance(wert, (int, float)):
        return float(wert)
    text = str(wert).strip().replace("€", "").replace(" ", "")
    if not text:
        return None
    if "," in text:
        text = text.replace(".", "").replace(",", ".")
    try:
        return float(text)
    except ValueError:
        return None


def wahr(wert):
    if isinstance(wert, str):
        return wert.strip().lower() in ("true", "1", "ja", "yes", "y")
    return bool(wert)


def norm_sku(sku):
    return "".join(str(sku).split()).upper()


def norm_ean(ean):
    return re.sub(r"\D", "", str(ean)).lstrip("0")


def text(wert):
    """Zelleninhalt als bereinigter Text ('nan' -> '', '65903.0' -> '65903')."""
    if wert is None:
        return ""
    t = str(wert).strip()
    if t.lower() in ("nan", "none", "nat"):
        return ""
    if re.fullmatch(r"\d+\.0", t):
        t = t[:-2]
    return t


# --------------------------------------------------------------------------- #
# Konfiguration & Eingabe
# --------------------------------------------------------------------------- #
def lade_config(pfad: Path) -> dict:
    with open(pfad, "rb") as f:
        return tomllib.load(f)


@dataclass
class Position:
    sku: str
    menge: float
    ean: str = ""
    hersteller: str = ""
    artikelnr: str = ""
    bezeichnung: str = ""
    letzter_ek: float | None = None

    @property
    def anzeige(self):
        return self.sku or f"EAN {self.ean}"


def lies_tabelle(pfad: Path, blatt) -> pd.DataFrame:
    if pfad.suffix.lower() == ".csv":
        for enc in ("utf-8-sig", "cp1252"):
            try:
                return pd.read_csv(pfad, sep=None, engine="python", dtype=str, encoding=enc)
            except UnicodeDecodeError:
                continue
        raise ValueError(f"CSV-Kodierung nicht erkannt: {pfad}")
    # xlrd (alte .xls, z. B. Mention-Export) schreibt harmlose Warnungen nach stdout
    with contextlib.redirect_stdout(io.StringIO()):
        return pd.read_excel(pfad, sheet_name=blatt or 0, dtype=str)


def lies_bedarfsliste(pfad: Path, cfg: dict) -> list[Position]:
    e = cfg["eingabe"]
    df = lies_tabelle(pfad, e.get("blatt"))
    spalten = {str(c).strip().lower(): c for c in df.columns}

    def spalte(schluessel, pflicht=False):
        name = (e.get(schluessel) or "").strip().lower()
        if not name:
            return None
        if name not in spalten:
            if pflicht:
                raise ValueError(f"Spalte '{e[schluessel]}' nicht in der Bedarfsliste. Vorhanden: "
                                 f"{list(df.columns)} – bitte config.toml [eingabe] prüfen.")
            log.warning("Spalte '%s' (%s) nicht gefunden – wird ignoriert", e[schluessel], schluessel)
            return None
        return spalten[name]

    sp_sku = spalte("spalte_herstellernummer", pflicht=True)
    sp_menge = spalte("spalte_menge", pflicht=True)
    sp_ean = spalte("spalte_ean")
    sp_hersteller = spalte("spalte_hersteller")
    sp_artikelnr = spalte("spalte_artikelnr")
    sp_bez = spalte("spalte_bezeichnung")
    sp_ek = spalte("spalte_letzter_ek")

    def wert(zeile, sp):
        return text(zeile[sp]) if sp is not None else ""

    positionen: dict[str, Position] = {}
    for nr, zeile in df.iterrows():
        pos = Position(sku=wert(zeile, sp_sku), menge=zahl(wert(zeile, sp_menge)) or 0.0,
                       ean=wert(zeile, sp_ean), hersteller=wert(zeile, sp_hersteller),
                       artikelnr=wert(zeile, sp_artikelnr), bezeichnung=wert(zeile, sp_bez),
                       letzter_ek=zahl(wert(zeile, sp_ek)) if sp_ek is not None else None)
        if not pos.sku and not pos.ean:
            if pos.artikelnr or pos.bezeichnung:
                log.warning("Zeile %s (%s): keine Herstellernummer/EAN – übersprungen",
                            nr + 2, pos.artikelnr or pos.bezeichnung)
            continue
        if pos.menge <= 0:
            log.info("Zeile %s (%s): Bedarf %g – übersprungen", nr + 2, pos.anzeige, pos.menge)
            continue
        if e.get("mengen_zusammenfassen", True):
            schluessel = norm_sku(pos.sku) if pos.sku else "EAN:" + norm_ean(pos.ean)
        else:
            schluessel = str(nr)
        if schluessel in positionen:
            positionen[schluessel].menge += pos.menge
        else:
            positionen[schluessel] = pos
    return list(positionen.values())


# --------------------------------------------------------------------------- #
# ITscope-API
# --------------------------------------------------------------------------- #
class ITscopeClient:
    def __init__(self, cfg: dict, account: str, api_key: str):
        a = cfg["api"]
        self.basis = a["basis_url"].rstrip("/")
        self.pfade = {"sku": a["such_pfad"], "ean": a.get("such_pfad_ean", "")}
        self.timeout = a.get("timeout_sekunden", 30)
        self.abstand = float(a.get("min_abstand_sekunden", 1.0))
        self.max_versuche = int(a.get("max_versuche", 5))
        self.session = requests.Session()
        self.session.auth = (account, api_key)
        self.session.headers.update({"Accept": "application/json",
                                     "User-Agent": "Einkauf-Skript/1.0"})
        self._letzte = 0.0

    def _warte_rate_limit(self):
        rest = self.abstand - (time.monotonic() - self._letzte)
        if rest > 0:
            time.sleep(rest)
        self._letzte = time.monotonic()

    def suche(self, art: str, wert: str):
        """art = 'sku' oder 'ean'. Antwort-JSON oder None (nicht gefunden).
        Wirft bei dauerhaften Fehlern."""
        if not self.pfade.get(art):
            return None
        url = self.basis + self.pfade[art].format(**{art: quote(wert, safe="")})
        for versuch in range(1, self.max_versuche + 1):
            self._warte_rate_limit()
            try:
                r = self.session.get(url, timeout=self.timeout)
            except requests.RequestException as ex:
                wartezeit = 2 ** versuch
                log.warning("%s: Netzwerkfehler (%s), neuer Versuch in %ss", wert, ex, wartezeit)
                time.sleep(wartezeit)
                continue

            if r.status_code == 404:
                return None
            if r.status_code == 429 or r.status_code >= 500:
                wartezeit = _retry_after(r) or 2 ** versuch
                log.warning("%s: HTTP %s, warte %ss (Versuch %s/%s)",
                            wert, r.status_code, wartezeit, versuch, self.max_versuche)
                time.sleep(wartezeit)
                continue
            if r.status_code in (401, 403):
                raise PermissionError(
                    f"HTTP {r.status_code} ({'Login abgelehnt' if r.status_code == 401 else 'Zugriff verweigert'})"
                    f" für {r.url} – Antwort von ITscope: {r.text.strip()[:500] or '(leer)'}")
            if not r.ok:
                raise RuntimeError(f"HTTP {r.status_code}: {r.text[:200]}")
            if not r.content.strip():
                return None
            return r.json()
        raise RuntimeError(f"keine Antwort nach {self.max_versuche} Versuchen")


def _retry_after(r):
    try:
        return max(1, int(float(r.headers.get("Retry-After", ""))))
    except ValueError:
        return None


class MockClient:
    """Liest {Herstellernummer}.json bzw. {EAN}.json (URL-kodiert, z. B. / -> %2F)
    aus einem Ordner – Test ohne API."""

    def __init__(self, ordner: Path):
        self.ordner = ordner

    def suche(self, art, wert):
        datei = self.ordner / f"{quote(wert, safe='')}.json"
        if not datei.exists():
            return None
        return json.loads(datei.read_text(encoding="utf-8"))


# --------------------------------------------------------------------------- #
# Auswertung
# --------------------------------------------------------------------------- #
@dataclass
class Angebot:
    lieferant: str
    lieferant_id: str
    preis: float
    bestand: float


@dataclass
class Ergebnis:
    pos: Position
    status: str = ""
    bezeichnung: str = ""
    puid: str = ""
    best: Angebot | None = None
    zweit: Angebot | None = None
    hinweis: str = ""
    angebote_gesamt: int = 0
    angebote_meine: list = field(default_factory=list)

    @property
    def sku(self):
        return self.pos.anzeige

    @property
    def menge(self):
        return self.pos.menge


def waehle_produkt(antwort, pos: Position, f):
    """Passendstes Produkt aus der Antwort: Herstellernummer/EAN exakt, Hersteller als Bonus."""
    if not antwort:
        return None
    produkte = als_liste(hole(antwort, f["produktliste"]))
    if not produkte and isinstance(antwort, dict) and hole(antwort, f["angebote"]) is not None:
        produkte = [antwort]          # Antwort ist direkt ein einzelnes Produkt
    produkte = [p for p in produkte if isinstance(p, dict)]
    if not produkte:
        return None

    def punkte(p):
        n = 0
        if pos.sku and norm_sku(hole(p, f["produkt_herstellernummer"]) or "") == norm_sku(pos.sku):
            n += 4
        if pos.ean and norm_ean(pos.ean) in {norm_ean(x) for x in als_liste(hole(p, f["produkt_ean"]))}:
            n += 4
        h = str(hole(p, f["produkt_hersteller"]) or "").lower()
        if pos.hersteller and h and (pos.hersteller.lower() in h or h in pos.hersteller.lower()):
            n += 1
        return n

    return max(produkte, key=punkte)   # bei Gleichstand der erste Treffer


def ist_meiner(roh: dict, a: Angebot, cfg_l: dict) -> bool:
    feld = (cfg_l.get("eigene_konditionen_feld") or "").strip()
    if feld and not wahr(hole(roh, [feld])):
        return False
    meine = {str(m).strip().lower() for m in cfg_l.get("meine", []) if str(m).strip()}
    if meine:
        return a.lieferant.strip().lower() in meine or a.lieferant_id.strip().lower() in meine
    return True


def bewerte(pos: Position, client, cfg: dict) -> Ergebnis:
    f = cfg["felder"]
    menge = pos.menge
    erg = Ergebnis(pos=pos, bezeichnung=pos.bezeichnung)
    produkt = waehle_produkt(client.suche("sku", pos.sku), pos, f) if pos.sku else None
    if produkt is None and pos.ean:
        produkt = waehle_produkt(client.suche("ean", pos.ean), pos, f)
        if produkt is not None:
            erg.hinweis = "über EAN gefunden"
    if produkt is None:
        erg.status = NICHT_GEFUNDEN
        erg.hinweis = "Keine Treffer bei ITscope" + (" (Herstellernr. und EAN)" if pos.sku and pos.ean else "")
        return erg

    erg.puid = str(hole(produkt, f["produkt_id"]) or "")
    erg.bezeichnung = str(hole(produkt, f["produkt_bezeichnung"]) or "") or pos.bezeichnung
    gefundene_sku = hole(produkt, f["produkt_herstellernummer"])
    if pos.sku and gefundene_sku and norm_sku(gefundene_sku) != norm_sku(pos.sku):
        erg.hinweis = (erg.hinweis + "; " if erg.hinweis else "") + \
            f"Abweichende Herstellernummer im Treffer: {gefundene_sku} – bitte prüfen"

    angebote = []
    roh_liste = als_liste(hole(produkt, f["angebote"]))
    erg.angebote_gesamt = len(roh_liste)
    for roh in roh_liste:
        if not isinstance(roh, dict):
            continue
        preis = zahl(hole(roh, f["preis"]))
        if preis is None or preis <= 0:
            continue
        a = Angebot(lieferant=str(hole(roh, f["lieferant_name"]) or "?"),
                    lieferant_id=str(hole(roh, f["lieferant_id"]) or ""),
                    preis=preis,
                    bestand=zahl(hole(roh, f["lagerbestand"])) or 0.0)
        if ist_meiner(roh, a, cfg["lieferanten"]):
            angebote.append(a)

    if not angebote:
        erg.status = KEIN_ANGEBOT
        erg.hinweis = (erg.hinweis + "; " if erg.hinweis else "") + \
            f"{erg.angebote_gesamt} Angebot(e) gesamt, keins von meinen Lieferanten mit Preis"
        return erg

    # Reihenfolge: genug Bestand < Teilbestand < kein Bestand, jeweils nach Preis
    def stufe(a):
        return 0 if a.bestand >= menge else (1 if a.bestand > 0 else 2)

    angebote.sort(key=lambda a: (stufe(a), a.preis))
    erg.angebote_meine = angebote
    erg.best = angebote[0]
    erg.zweit = next((a for a in angebote[1:] if a.lieferant != erg.best.lieferant), None)
    erg.status = {0: OK, 1: TEILBESTAND, 2: KEIN_BESTAND}[stufe(erg.best)]
    if erg.status == TEILBESTAND:
        erg.hinweis = (erg.hinweis + "; " if erg.hinweis else "") + \
            f"Nur {erg.best.bestand:g} von {menge:g} verfügbar"
    elif erg.status == KEIN_BESTAND:
        erg.hinweis = (erg.hinweis + "; " if erg.hinweis else "") + "Kein Lieferant hat Bestand"
    return erg


# --------------------------------------------------------------------------- #
# Excel-Ausgabe
# --------------------------------------------------------------------------- #
SPALTEN = ["Herstellernummer", "Bezeichnung", "Menge", "Bester Lieferant", "Preis netto",
           "Lagerbestand", "Zweitbester Lieferant", "Preis netto (2.)", "Gesamtsumme",
           "ITscope-Link", "Status", "Hinweis", "Eigene Art.-Nr.", "Letzter EK", "EAN"]
BREITEN = [20, 50, 8, 24, 12, 12, 24, 14, 14, 16, 26, 50, 16, 12, 16]
EURO = '#,##0.00 "€"'


def schreibe_excel(ergebnisse: list[Ergebnis], ziel: Path, cfg: dict):
    wb = Workbook()
    ws = wb.active
    ws.title = "Einkauf"
    ws.append(SPALTEN)
    for zelle in ws[1]:
        zelle.font = Font(bold=True)
        zelle.fill = KOPF
        zelle.alignment = Alignment(vertical="center", wrap_text=True)
    link_vorlage = cfg["api"].get("produkt_link", "")

    for erg in ergebnisse:
        best, zweit = erg.best, erg.zweit
        ws.append([erg.sku, erg.bezeichnung, erg.menge,
                   best.lieferant if best else None, best.preis if best else None,
                   best.bestand if best else None,
                   zweit.lieferant if zweit else None, zweit.preis if zweit else None,
                   None, None, erg.status, erg.hinweis,
                   erg.pos.artikelnr, erg.pos.letzter_ek, erg.pos.ean])
        r = ws.max_row
        if best:
            ws.cell(r, 9, f"=C{r}*E{r}")
        if erg.puid and link_vorlage:
            c = ws.cell(r, 10, "Produkt öffnen")
            c.hyperlink = link_vorlage.format(puid=erg.puid, sku=quote(erg.pos.sku, safe=""))
            c.font = Font(color="0563C1", underline="single")
        for sp in (5, 8, 9, 14):
            ws.cell(r, sp).number_format = EURO
        if erg.status in FARBE:
            for sp in range(1, len(SPALTEN) + 1):
                if sp != 10:
                    ws.cell(r, sp).fill = FARBE[erg.status]

    letzte = ws.max_row
    ws.append([])
    ws.append([None] * 7 + ["Summe:", f"=SUM(I2:I{letzte})"])
    ws.cell(ws.max_row, 8).font = Font(bold=True)
    ws.cell(ws.max_row, 9).font = Font(bold=True)
    ws.cell(ws.max_row, 9).number_format = EURO

    probleme = [e for e in ergebnisse if e.status in (NICHT_GEFUNDEN, FEHLER, KEIN_ANGEBOT)]
    if probleme:
        ws.append([])
        ws.append([f"Nicht gefunden / nicht bestellbar ({len(probleme)}):"])
        ws.cell(ws.max_row, 1).font = Font(bold=True, color="C00000")
        for e in probleme:
            ws.append([e.sku, e.bezeichnung, e.menge, None, None, None, None, None, None, None,
                       e.status, e.hinweis, e.pos.artikelnr, None, e.pos.ean])
            for sp in range(1, len(SPALTEN) + 1):
                ws.cell(ws.max_row, sp).fill = FARBE[e.status]

    for i, b in enumerate(BREITEN, 1):
        ws.column_dimensions[get_column_letter(i)].width = b
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(SPALTEN))}{letzte}"

    # Zweites Blatt: alle Angebote meiner Lieferanten zum Nachvollziehen
    wa = wb.create_sheet("Alle Angebote")
    wa.append(["Herstellernummer", "Lieferant", "Lieferanten-ID", "Preis netto", "Lagerbestand",
               "Reicht für Menge"])
    for zelle in wa[1]:
        zelle.font = Font(bold=True)
        zelle.fill = KOPF
    for e in ergebnisse:
        for a in e.angebote_meine:
            wa.append([e.sku, a.lieferant, a.lieferant_id, a.preis, a.bestand,
                       "ja" if a.bestand >= e.menge else "nein"])
            wa.cell(wa.max_row, 4).number_format = EURO
    for i, b in enumerate([20, 28, 16, 14, 14, 16], 1):
        wa.column_dimensions[get_column_letter(i)].width = b

    ziel.parent.mkdir(parents=True, exist_ok=True)
    wb.save(ziel)


# --------------------------------------------------------------------------- #
# Hauptprogramm
# --------------------------------------------------------------------------- #
def richte_logging_ein():
    ordner = BASIS / "logs"
    ordner.mkdir(exist_ok=True)
    datei = ordner / f"einkauf_{datetime.now():%Y-%m-%d}.log"
    logging.basicConfig(
        level=logging.INFO,
        format="%(asctime)s %(levelname)-7s %(message)s",
        handlers=[logging.FileHandler(datei, encoding="utf-8"), logging.StreamHandler(sys.stdout)],
    )


def finde_eingangsdateien(cfg: dict) -> tuple[Path, list[Path]]:
    ordner = Path(cfg["eingabe"].get("ordner") or "eingang")
    if not ordner.is_absolute():
        ordner = BASIS / ordner
    ordner.mkdir(parents=True, exist_ok=True)
    dateien = [d for d in ordner.iterdir()
               if d.is_file() and d.suffix.lower() in (".xls", ".xlsx", ".csv")
               and not d.name.startswith("~$")]
    return ordner, sorted(dateien, key=lambda d: d.stat().st_mtime)


def zielpfad(cfg: dict, eingabe: Path) -> Path:
    a = cfg["ausgabe"]
    ordner = Path(a["ordner"])
    if not ordner.is_absolute():
        ordner = BASIS / ordner
    ziel = ordner / a["dateiname"].format(datum=f"{datetime.now():%Y-%m-%d_%H%M}", name=eingabe.stem)
    n = 2
    while ziel.exists():
        ziel = ziel.with_name(f"{ziel.stem.rsplit('#', 1)[0]}#{n}{ziel.suffix}")
        n += 1
    return ziel


def verarbeite(eingabe: Path, client, cfg: dict, ziel: Path) -> Path:
    positionen = lies_bedarfsliste(eingabe, cfg)
    log.info("%d Positionen mit Bedarf aus %s gelesen", len(positionen), eingabe.name)

    ergebnisse = []
    for i, pos in enumerate(positionen, 1):
        try:
            erg = bewerte(pos, client, cfg)
        except PermissionError:
            raise                      # Zugangsdaten falsch – betrifft alle Artikel
        except Exception as ex:
            log.exception("%s: Fehler bei der Verarbeitung", pos.anzeige)
            erg = Ergebnis(pos=pos, bezeichnung=pos.bezeichnung, status=FEHLER, hinweis=str(ex)[:300])
        ergebnisse.append(erg)
        log.info("[%d/%d] %s x%g -> %s%s", i, len(positionen), pos.anzeige, pos.menge, erg.status,
                 f" ({erg.best.lieferant}, {erg.best.preis:.2f} €)" if erg.best else "")

    try:
        schreibe_excel(ergebnisse, ziel, cfg)
    except PermissionError:
        ziel = ziel.with_name(ziel.stem + f"_{datetime.now():%H%M%S}" + ziel.suffix)
        log.warning("Zieldatei gesperrt (in Excel geöffnet?) – schreibe nach %s", ziel.name)
        schreibe_excel(ergebnisse, ziel, cfg)

    zaehler = {}
    for e in ergebnisse:
        zaehler[e.status] = zaehler.get(e.status, 0) + 1
    log.info("Fertig: %s", ", ".join(f"{k}: {v}" for k, v in zaehler.items()) or "keine Positionen")
    log.info("Ergebnis: %s", ziel)
    nicht = [e.sku for e in ergebnisse if e.status in (NICHT_GEFUNDEN, FEHLER)]
    if nicht:
        log.warning("Nicht gefunden/Fehler: %s", ", ".join(nicht))
    return ziel


def main(argv=None) -> int:
    p = argparse.ArgumentParser(description="Täglicher Einkauf über die ITscope-API")
    p.add_argument("bedarfsliste", nargs="*",
                   help="Excel/CSV-Datei(en). Ohne Angabe: alle Dateien im Eingangsordner")
    p.add_argument("--config", default=str(BASIS / "config.toml"))
    p.add_argument("--ausgabe", help="Pfad der Ergebnisdatei (.xlsx), nur bei einer Eingabedatei")
    p.add_argument("--mock", help="Ordner mit {Herstellernummer}.json statt echter API (Test)")
    p.add_argument("--oeffnen", action="store_true", help="Ergebnis danach in Excel öffnen")
    args = p.parse_args(argv)

    richte_logging_ein()
    try:
        cfg = lade_config(Path(args.config))
    except (OSError, tomllib.TOMLDecodeError) as ex:
        log.error("Konfiguration nicht lesbar: %s", ex)
        return 2

    aus_eingang = not args.bedarfsliste
    if aus_eingang:
        ordner, dateien = finde_eingangsdateien(cfg)
        if not dateien:
            log.info("Keine Bedarfsliste in %s – nichts zu tun.", ordner)
            if args.oeffnen and hasattr(os, "startfile"):
                print("\n>>> Der Ordner 'eingang' ist leer. Bitte die Bedarfsliste dort hineinlegen"
                      "\n>>> (er wird jetzt geöffnet) oder die Datei direkt auf Einkauf_jetzt.bat ziehen.")
                os.startfile(ordner)
            return 0
    else:
        dateien = [Path(d).resolve() for d in args.bedarfsliste]

    if args.mock:
        client = MockClient(Path(args.mock))
    else:
        load_dotenv(BASIS / ".env")
        account = os.getenv("ITSCOPE_ACCOUNT_ID", "").strip()
        api_key = os.getenv("ITSCOPE_API_KEY", "").strip()
        if not account or not api_key:
            log.error("ITSCOPE_ACCOUNT_ID und ITSCOPE_API_KEY müssen in .env gesetzt sein")
            return 2
        client = ITscopeClient(cfg, account, api_key)

    l = cfg["lieferanten"]
    if not l.get("meine") and not l.get("eigene_konditionen_feld"):
        log.warning("Keine 'meinen Lieferanten' konfiguriert – es werden ALLE Angebote berücksichtigt!")

    rueckgabe = 0
    for eingabe in dateien:
        ziel = Path(args.ausgabe) if (args.ausgabe and len(dateien) == 1) else zielpfad(cfg, eingabe)
        try:
            ziel = verarbeite(eingabe, client, cfg, ziel)
        except PermissionError as ex:
            log.error("Abbruch: %s", ex)
            return 3
        except Exception as ex:
            log.error("Bedarfsliste %s nicht verarbeitbar: %s", eingabe.name, ex)
            rueckgabe = 2
            continue
        if aus_eingang and cfg["eingabe"].get("erledigt_verschieben", True):
            archiv = eingabe.parent / "erledigt"
            archiv.mkdir(exist_ok=True)
            neu = archiv / f"{datetime.now():%Y-%m-%d_%H%M}_{eingabe.name}"
            try:
                shutil.move(eingabe, neu)
                log.info("Bedarfsliste verschoben nach %s", neu)
            except OSError as ex:
                log.warning("Konnte %s nicht verschieben (%s)", eingabe.name, ex)
        if args.oeffnen and hasattr(os, "startfile"):
            os.startfile(ziel)
    return rueckgabe


if __name__ == "__main__":
    sys.exit(main())
